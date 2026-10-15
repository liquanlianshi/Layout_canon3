# -*- coding: utf-8 -*-
"""Structure checks for the generated Word narration script.

The .docx is produced by scripts/make_speech_doc.py; these tests read it back
with python-docx and assert that it really is a usable 4-minute storyboard.
"""
import importlib.util
import os

import pytest

docx = pytest.importorskip("docx")

from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC = os.path.join(ROOT, "视频讲解文稿-4分钟.docx")

SEGMENTS = 10
MIN_CHARS = 850          # ~3.5 minutes at 245 chars/min
MAX_CHARS = 1080         # ~4.4 minutes at 245 chars/min


@pytest.fixture(scope="module")
def document():
    """Regenerate the document first, so these tests cover the *script* and
    not merely whatever .docx happens to sit on disk."""
    spec = importlib.util.spec_from_file_location(
        "make_speech_doc", os.path.join(ROOT, "scripts", "make_speech_doc.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.main()
    assert os.path.exists(DOC), "script did not produce the expected file"
    return docx.Document(DOC)


def test_document_exists_and_is_a_docx(document):
    assert len(document.paragraphs) > 10


def test_table_uses_a_single_fixed_layout(document):
    """Exactly one <w:tblLayout w:type="fixed"/> may be present: a duplicate
    would be invalid OOXML and Word can refuse to open the file."""
    tbl = document.tables[0]._tbl
    layouts = tbl.tblPr.findall(qn("w:tblLayout"))
    assert len(layouts) == 1
    assert layouts[0].get(qn("w:type")) == "fixed"


def test_title_and_meta_present(document):
    text = "\n".join(p.text for p in document.paragraphs)
    assert "新质生产力" in text
    assert "分镜" in text
    assert "4 分 00 秒" in text
    assert "字/分钟" in text


def test_storyboard_table_has_one_row_per_segment(document):
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.columns) == 5
    assert len(table.rows) == SEGMENTS + 1          # header + segments
    header = [c.text for c in table.rows[0].cells]
    assert header[0] == "时间码"
    assert header[2] == "画面 / 操作"
    assert header[3] == "中文讲解词"


def test_every_segment_has_time_scene_and_narration(document):
    table = document.tables[0]
    for row in table.rows[1:]:
        time_code, dur, scene, zh, en = (c.text.strip() for c in row.cells)
        assert ":" in time_code and time_code[0].isdigit(), time_code
        assert dur.endswith("s")
        assert len(scene) > 10, time_code
        assert len(zh) > 30, time_code
        assert en and en[0].isascii(), time_code


def test_timeline_covers_zero_to_four_minutes(document):
    table = document.tables[0]
    codes = [row.cells[0].text.strip() for row in table.rows[1:]]
    assert codes[0].startswith("0:00")
    assert codes[-1].endswith("4:00")
    total = sum(int(row.cells[1].text.strip().rstrip("s")) for row in table.rows[1:])
    assert total == 240, total                      # exactly four minutes


def test_narration_length_matches_a_four_minute_video(document):
    table = document.tables[0]
    zh_total = sum(len(row.cells[3].text.strip()) for row in table.rows[1:])
    assert MIN_CHARS <= zh_total <= MAX_CHARS, zh_total


def test_continuous_read_through_section_present(document):
    text = "\n".join(p.text for p in document.paragraphs)
    assert "连读版" in text
    assert "录制与演示提示" in text
    # the read-through repeats all ten narration blocks
    for marker in ("1.（0:00", "10.（3:50"):
        assert marker in text, marker
