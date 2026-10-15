# -*- coding: utf-8 -*-
"""Generate the 4-minute video narration script as a Word (.docx) document.

Output: 视频讲解文稿-4分钟.docx  (project root)

The script is deterministic: re-running it overwrites the document.
Timing is planned for a Chinese narration pace of roughly 230 characters per
minute, which maps the 10 segments onto 4 minutes (240 s).
"""
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "视频讲解文稿-4分钟.docx")

CJK = "微软雅黑"
LATIN = "Calibri"
INK = RGBColor(0x33, 0x30, 0x2C)
MUTED = RGBColor(0x6E, 0x67, 0x5F)
ACCENT = RGBColor(0xC4, 0x84, 0x6F)

# time, duration, scene / actions, Chinese narration, English cue
SEGMENTS = [
    ("0:00–0:20", "20s",
     "浏览器打开首页，页面停在顶部，缓慢下移一小段，展示玻璃拟态卡片与彩色光斑背景。",
     "大家好，这是我做的数据可视化网站——新质生产力与社会发展。整站没有一张写死的图形，每条曲线、每个气泡都是浏览器实时算出来的。接下来四分钟，我带您看它如何用图形讲清两者的关系。",
     "Intro: a fully data-driven, bilingual visualisation site."),

    ("0:20–0:47", "27s",
     "鼠标依次掠过四张 KPI 卡片，再向下滚到“全部指标速览”网格，指一下其中两张小卡。",
     "顶部四张卡片是 2024 年的关键读数：劳动生产率、数字经济占比、居民收入持续上升，单位 GDP 能耗逐年下降，百分比是相对 2015 年的变化。再往下是全部十一项指标的速览，每张小卡都带一条十年趋势，绿色代表改善，橙色代表需关注。",
     "KPIs plus an all-indicator grid with ten-year sparklines."),

    ("0:47–1:22", "35s",
     "停在“新质生产力指数 × 社会发展指数”，鼠标沿两条曲线走一遍；再指左下角的指数仪表，然后移到右侧桑基图并悬停一条连线，展示相关系数。",
     "这一张是全站核心：把五项新质生产力指标标准化后合成一条指数，社会发展指标合成另一条，两条曲线十年同步抬升。左下角是 2024 年的指数仪表；右侧这张桑基图更进一步——连线的粗细不是画出来的，而是用面板数据算出的皮尔逊相关系数，只保留正相关，可以看到研发投入与数字技术，主要流向生产率和居民收入。",
     "Composite index, the 2024 gauge, and a sankey whose widths are real correlations."),

    ("1:22–1:57", "35s",
     "依次展示：气泡散点（悬停一个气泡）、地区×指标热力图、按年播放的排名动画、11×11 相关矩阵。",
     "继续往下看。气泡散点把十年、十个地区的数据全铺开：横轴是研发投入强度，纵轴是劳动生产率，气泡大小是数字经济占比，虚线是自动拟合的回归线，正向关系非常明显。热力图对比各地区在多维指标上的得分；这一段动画按年份播放各地区数字经济占比的排名变化；最后是十一乘十一的相关矩阵，冷色是负相关，暖色是正相关。",
     "Scatter with fitted trend line, heatmap, animated ranking, correlation matrix."),

    ("1:57–2:22", "25s",
     "展示单位 GDP 能耗曲线、就业×高新技术企业双轴图、2015 vs 2024 对比柱状图。",
     "这三张图回答“社会发展得到了什么”：能耗强度十年持续下降，三条地区曲线都向下；柱线双轴把高新技术企业数和城镇新增就业放在一起看；最后一组柱状图把 2015 与 2024 归一化对比，增长倍数直接标在柱顶。",
     "Outcomes: energy intensity, jobs vs high-tech firms, 2015 vs 2024."),

    ("2:22–2:37", "15s",
     "点击右上角切换到“数据分析工作台”，页面滚到顶部。",
     "不过图形只是结果。真正做分析的地方，是第二个页面——数据分析工作台。它把 pandas 的常用操作，全部做成了可以点击的按钮。",
     "Switch to the analysis workbench."),

    ("2:37–3:07", "30s",
     "展示左侧数据源面板；在筛选框输入 indicator == 'digital_share' and year >= 2020 并点“应用筛选”；表格与图表刷新；指一下操作记录里的 pandas 代码；再点一个列标签。",
     "左边是数据源，默认加载一千一百行长表，也能拖入自己的 CSV，或直接粘贴文本。这里是筛选：输入 indicator 等于 digital_share、且年份大于等于 2020，点应用——表格立刻缩到五十行，图表同步重画，下面还记下了等价的 pandas 代码。排序与选择列也是同样的即时联动。",
     "Load, filter, sort, project columns — every step logs its pandas equivalent."),

    ("3:07–3:35", "28s",
     "在“分组与聚合”里多选年份、选均值并执行；再用透视表把地区展开；最后切换一次图表类型。",
     "分组聚合支持多选字段和八种聚合函数，我按年份求均值，图立刻变成十年趋势。透视表选好行、列和数值，一次就能把地区展开成十条系列。右上角还可以切换折线、柱状、堆积面积、环形、散点和箱线六种图表类型，所有图形都跟着当前数据走。",
     "Group-by, pivot table and six chart types."),

    ("3:35–3:50", "15s",
     "依次点击“统计描述”“相关性”“分布直方图”，再指向下方的数据健康面板。",
     "再点几下，就能得到描述统计、相关矩阵、直方图和频次统计；这张数据健康面板会给出每列的完整度、唯一值与重复行情况。分析做完，一键导出结果。",
     "describe, corr, hist, value_counts, plus the data-health panel and CSV export."),

    ("3:50–4:00", "10s",
     "点击右上角 EN 切换英文，页面文案与图表标题整体切换；最后回到首页顶部。",
     "最后，整站中英双语，点右上角即可切换，内容与图形完全一致。这就是我的作品：让数据自己说话。谢谢观看。",
     "Bilingual switch and closing."),
]


def style_run(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = LATIN
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    return run


def heading(doc, text, size=16, color=INK, space_before=10, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    style_run(p.add_run(text), size=size, bold=True, color=color)
    return p


def body(doc, text, size=10.5, color=INK, indent=0, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    style_run(p.add_run(text), size=size, color=color)
    return p


def main():
    total_chars = sum(len(s[3]) for s in SEGMENTS)

    doc = Document()

    # ---- page setup: A4 landscape so the storyboard table fits -----------
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    # ---- title block -----------------------------------------------------
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(t.add_run("新质生产力 × 社会发展 · 数据可视化网站"), size=20, bold=True, color=INK)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(s.add_run("四分钟演示视频 · 讲解文稿（分镜 + 口播）"), size=13, bold=True, color=ACCENT)

    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pace = int(round(total_chars / 4.0))
    style_run(m.add_run(
        "适用时长 4 分 00 秒 · 共 10 个分镜 · 口播约 %d 字 · 语速约 %d 字/分钟" % (total_chars, pace)
    ), size=10, color=MUTED)

    # ---- storyboard table ------------------------------------------------
    heading(doc, "一、分镜与口播（照此录制）", size=14)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.autofit = False          # also emits <w:tblLayout w:type="fixed"/>
    widths = [Cm(1.9), Cm(1.3), Cm(4.8), Cm(11.2), Cm(4.2)]
    headers = ["时间码", "时长", "画面 / 操作", "中文讲解词", "English cue"]
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(h), size=10, bold=True, color=INK)

    for time_code, dur, scene, zh, en in SEGMENTS:
        cells = table.add_row().cells
        for i, value in enumerate((time_code, dur, scene, zh, en)):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            if i < 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(p.add_run(value), size=9.5,
                      bold=(i == 0), color=MUTED if i == 4 else INK)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]

    # ---- continuous read-through ----------------------------------------
    doc.add_page_break()
    heading(doc, "二、连读版（计时练习用）", size=14)
    body(doc, "按顺序一口气念完下面十段，即为完整四分钟口播。建议先按段掐表，再整体连读两遍。",
         size=10, color=MUTED, space_after=8)
    for idx, (time_code, dur, _scene, zh, _en) in enumerate(SEGMENTS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.35
        style_run(p.add_run("%d.（%s，%s）" % (idx, time_code, dur)), size=10, bold=True, color=ACCENT)
        style_run(p.add_run(zh), size=11, color=INK)

    # ---- recording notes -------------------------------------------------
    heading(doc, "三、录制与演示提示", size=14)
    for tip in [
        "浏览器建议使用 Chrome 或 Edge，窗口宽度 1440px 以上，缩放 100%，以获得最佳玻璃拟态与图表效果。",
        "录制前先打开页面静置到动画结束；按段录制，方便后期剪辑与重录单段。",
        "鼠标移动放慢，悬停气泡、桑基连线、热力图色块各停留 1 秒，让相关系数与数值提示可见。",
        "讲解时语言可随时切换：右上角“中文 / EN”按钮，切到英文可验证双语内容一致。",
        "若需英文版录制，可按表格最后一列的 English cue 逐段展开，内容与中文一一对应。",
        "全部数据为教学演示用合成数据，页面页脚与说明文档中均已标注，讲解时可主动说明。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        style_run(p.add_run(tip), size=10.5, color=INK)

    # ---- closing footnote ------------------------------------------------
    f = doc.add_paragraph()
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.paragraph_format.space_before = Pt(14)
    style_run(f.add_run("配套素材：index.html（网站入口）· README.md（说明文档）· data/（默认数据集）"),
              size=9.5, color=MUTED)

    doc.save(OUT)
    print("wrote speech script: %d segments, %d narration characters, %d bytes"
          % (len(SEGMENTS), total_chars, os.path.getsize(OUT)))
    return OUT


if __name__ == "__main__":
    main()
